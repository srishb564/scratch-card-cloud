from flask import Flask, render_template, request, redirect, url_for, jsonify
import uuid
import os
import random
from openpyxl import Workbook, load_workbook

app = Flask(__name__, template_folder="templates", static_folder="static")

EXCEL_FILE = "scratch_cards.xlsx"

REWARDS = ["₹50", "₹100", "Better Luck Next Time"]

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "ScratchCards"
        ws.append(["Card Name", "Card ID", "Reward", "Link", "Scratched"])
        wb.save(EXCEL_FILE)

init_excel()

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        card_name = request.form["card_name"]
        card_id = str(uuid.uuid4())
        reward = random.choice(REWARDS)
        link = request.host_url + "scratch/" + card_id

        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        ws.append([card_name, card_id, reward, link, "NO"])
        wb.save(EXCEL_FILE)

        return redirect(url_for("index"))

    return render_template("index.html")

@app.route("/scratch/<card_id>")
def scratch(card_id):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    reward = "Better Luck Next Time"

    for row in ws.iter_rows(min_row=2):
        if row[1].value == card_id:
            reward = row[2].value
            break

    return render_template("scratch.html", card_id=card_id, reward=reward)

@app.route("/mark_scratched", methods=["POST"])
def mark_scratched():
    card_id = request.json["card_id"]

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[1].value == card_id:
            row[4].value = "YES"
            break

    wb.save(EXCEL_FILE)
    return jsonify({"status": "updated"})

if __name__ == "__main__":
    app.run(debug=True)
